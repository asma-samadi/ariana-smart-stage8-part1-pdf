# Regenerates the fictional workbook. Run from this folder: python3 build_sample_workbook.py  (needs: pip install openpyxl)
# headers.json = output of: node -e "import('../backend/src/lib/sheetsRow.js').then(m=>console.log(JSON.stringify(m.SHEET_HEADERS)))"
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
headers = json.load(open('headers.json'))
wb = Workbook()
hdr_fill = PatternFill('solid', fgColor='1F4E79'); hdr_font = Font(bold=True, color='FFFFFF')

ws = wb.active; ws.title = 'README'
for r in [
 ['Ariana Smart - FICTIONAL SAMPLE DATA'],
 ['Educational demo. Not affiliated with Ariana Afghan Airlines. Every row below is invented.'],
 ['No real passengers, passport/ID numbers, document files, links, emails, phone numbers, tokens or credentials are in this workbook.'],
 [''],
 ['Sheets'],
 ['Applications', 'Same columns, in the same order, as the Google Sheet the backend writes to (backend/src/lib/sheetsRow.js).'],
 ['Routes', 'Fictional demo route network (same airports as the app).'],
 ['Summary', 'Live COUNTIF formulas over the Applications sheet.'],
 [''],
 ['Tip: to try the real integration, create your own empty Google Sheet and follow docs/SETUP_GOOGLE_SHEETS.md. Do not import real passenger data here.'],
]: ws.append(r)
ws['A1'].font = Font(bold=True, size=14); ws['A5'].font = Font(bold=True)
ws.column_dimensions['A'].width = 18; ws.column_dimensions['B'].width = 110

ws = wb.create_sheet('Applications'); ws.append(headers)
routes = [('KBL','DXB'),('KBL','IST'),('HEA','KBL'),('MZR','DEL'),('DXB','KBL'),('IST','KBL'),('DEL','KBL'),('KBL','MZR'),('KBL','HEA'),('DXB','IST')]
statuses = ['submitted','under_review','documents_required','approved_demo','rejected','cancelled','submitted','under_review','approved_demo','submitted']
nat = ['Testland','Exampleland','Sampleland','Testland','Demoria','Exampleland','Sampleland','Testland','Demoria','Testland']
for i in range(10):
    o,d = routes[i]; ret = f'2026-12-{10+i:02d}' if i % 3 == 0 else ''
    ws.append([f'ASM-SAMPLE-{1000+i}', f'2026-10-{1+i:02d}T09:{10+i}:00.000Z', statuses[i], o, d, 'return' if ret else 'one_way',
               f'2026-11-{1+i:02d}', ret, 1 + (i % 3), nat[i], f'2026-10-{1+i:02d}T09:{11+i}:30.000Z'])
for c in ws[1]: c.fill = hdr_fill; c.font = hdr_font; c.alignment = Alignment(wrap_text=True, vertical='center')
for col, w in zip('ABCDEFGHIJK', [20,26,20,9,12,11,15,13,11,14,26]): ws.column_dimensions[col].width = w
ws.freeze_panes = 'A2'
t = Table(displayName='Applications', ref=f'A1:K{ws.max_row}'); t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True); ws.add_table(t)

ws = wb.create_sheet('Routes'); ws.append(['Origin','Destination','Note'])
for o,d in routes[:6]: ws.append([o,d,'Fictional demo route'])
for c in ws[1]: c.fill = hdr_fill; c.font = hdr_font
ws.column_dimensions['C'].width = 24

ws = wb.create_sheet('Summary'); ws.append(['Status','Count'])
for s in ['submitted','under_review','documents_required','approved_demo','rejected','cancelled']:
    r = ws.max_row + 1; ws.append([s, f'=COUNTIF(Applications!$C$2:$C$1000,A{r})'])
r = ws.max_row + 1; ws.append(['Total', f'=SUM(B2:B{r-1})'])
for c in ws[1]: c.fill = hdr_fill; c.font = hdr_font
ws[f'A{r}'].font = Font(bold=True); ws[f'B{r}'].font = Font(bold=True)
ws.column_dimensions['A'].width = 22
wb.save('ariana-smart-sample-data.xlsx')
